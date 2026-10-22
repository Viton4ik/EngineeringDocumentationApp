"""Предпросмотр любого файла в браузере: конвертация в PDF/PNG/MP4 с кэшем, HTML-просмотр без LibreOffice."""
from __future__ import annotations
import io, re, csv, html, shutil, zipfile, hashlib, tempfile, threading, subprocess
from pathlib import Path
from datetime import datetime
from email import policy
from email.parser import BytesParser
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote
from flask import request
from .config import *  # noqa: F401,F403 - наборы расширений и пути к инструментам
from .config import Image
from .catalog import find_record, save_title, title_from_pdf
from .i18n import t

EXECUTOR = ThreadPoolExecutor(max_workers=2)
JOBS = {}
LO_LOCK = threading.Lock()

def cache_key(p: Path, target):
    st = p.stat()
    return hashlib.sha1(f"{p}|{st.st_mtime}|{st.st_size}|{target}".encode()).hexdigest()

def run(cmd, timeout):
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    r = subprocess.run(cmd, capture_output=True, timeout=timeout, creationflags=flags)
    if r.returncode != 0:
        raise RuntimeError((r.stderr or r.stdout or b"").decode(errors="replace")[-400:] or f"code {r.returncode}")
    return r

def convert_office(src: Path, out: Path):
    with LO_LOCK, tempfile.TemporaryDirectory(dir=DATA_DIR / "tmp") as tmp:
        profile = (DATA_DIR / "lo_profile").resolve().as_uri()
        run([SOFFICE, "--headless", "--norestore", "--nologo", f"-env:UserInstallation={profile}",
             "--convert-to", "pdf", "--outdir", tmp, str(src)], timeout=240)
        produced = Path(tmp) / (src.stem + ".pdf")
        if not produced.exists(): raise RuntimeError("LibreOffice: no PDF")
        shutil.move(str(produced), out)
    rec = find_record(src)
    if rec and rec["title_source"] == "file":  # старые .doc/.ppt: название берём из полученного PDF
        try: title = title_from_pdf(out)
        except Exception: title = ""
        if title: save_title(rec, title)

def convert_image(src: Path, out: Path):
    with Image.open(src) as im:
        im.seek(0)
        im.thumbnail((4000, 4000))
        im.convert("RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB").save(out, "PNG")

def convert_media(src: Path, out: Path):
    tmp = out.with_suffix(".part" + out.suffix)
    if out.suffix == ".mp4":
        args = ["-c:v", "libx264", "-preset", "veryfast", "-crf", "26", "-vf", "scale='min(1280,iw)':-2", "-c:a", "aac", "-movflags", "+faststart"]
    else:
        args = ["-vn", "-c:a", "libmp3lame", "-q:a", "4"]
    run([FFMPEG, "-y", "-loglevel", "error", "-i", str(src), *args, str(tmp)], timeout=3600)
    tmp.replace(out)

def converted(p: Path, target, kind, func, label_key):
    key = cache_key(p, target)
    out = CACHE_DIR / f"{key}.{target}"
    url = f"/api/cache/{out.name}/{quote(p.stem)}.{target}"  # имя файла в URL видно в окне просмотра PDF
    if out.exists(): return {"kind": kind, "url": url, "name": p.name}
    job = JOBS.get(key)
    if job is None or (job.done() and job.exception() and request.args.get("retry")):
        def task():
            try: func(p, out)
            except subprocess.TimeoutExpired: raise RuntimeError("timeout")
        job = JOBS[key] = EXECUTOR.submit(task)
    if job.done():
        if job.exception(): return {"kind": "error", "name": p.name, "message": t("pv_convert_failed", label=t(label_key), error=job.exception())}
        return {"kind": kind, "url": url, "name": p.name}
    return {"kind": "pending", "name": p.name, "message": t(label_key) + "… " + t("pv_first_time")}

# ---------- HTML-просмотр ----------
def h(s): return html.escape(str(s if s is not None else ""))

def docx_html(p):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(p))
    parts = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            if not para.text.strip(): continue
            style = (para.style.name if para.style is not None else "").lower()
            level = 1 if style.startswith(("title", "название")) else int(style[-1]) + 1 if style[:-1].strip() in ("heading", "заголовок") and style[-1].isdigit() else 0
            parts.append(f"<h{min(level,4)}>{h(para.text)}</h{min(level,4)}>" if level else f"<p>{h(para.text)}</p>")
        elif tag == "tbl":
            tbl = Table(child, d)
            parts.append("<table>" + "".join("<tr>" + "".join(f"<td>{h(c.text)}</td>" for c in row.cells) + "</tr>" for row in tbl.rows) + "</table>")
    return "".join(parts) or f"<p>{h(t('pv_empty'))}</p>"

def xlsx_html(p, max_rows=500, max_cols=40):
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets[:20]:
            rows = []
            for i, row in enumerate(ws.iter_rows(max_col=max_cols, values_only=True)):
                if i >= max_rows: rows.append(f"<tr><td colspan='{max_cols}'>… {h(t('pv_first_rows', n=max_rows))}</td></tr>"); break
                rows.append("<tr>" + "".join(f"<td>{h('' if v is None else v)}</td>" for v in row) + "</tr>")
            out.append(f"<h3>{h(t('pv_sheet'))}: {h(ws.title)}</h3><table>{''.join(rows)}</table>")
    finally:
        wb.close()
    return "".join(out)

def pptx_html(p):
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(f"<section class='slide'><div class='n'>{h(t('pv_slide'))} {i}</div>" + "".join(f"<p>{h(x)}</p>" for x in texts) + "</section>")
    return "".join(out) or f"<p>{h(t('pv_empty'))}</p>"

def csv_html(p, max_rows=2000):
    with p.open("rb") as f: raw = f.read(4 * 1024 * 1024)
    text = raw.decode("utf-8-sig", errors="replace")
    try: dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error: dialect = csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= max_rows: rows.append(f"<tr><td>… {h(t('pv_first_rows', n=max_rows))}</td></tr>"); break
        rows.append("<tr>" + "".join(f"<td>{h(c)}</td>" for c in row) + "</tr>")
    return f"<table>{''.join(rows)}</table>"

def eml_html(p):
    with p.open("rb") as f: msg = BytesParser(policy=policy.default).parse(f)
    head = "".join(f"<tr><td><b>{k}</b></td><td>{h(msg.get(k, ''))}</td></tr>" for k in ("From", "To", "Cc", "Date", "Subject"))
    body = msg.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    if body is not None and body.get_content_type() != "text/plain": content = re.sub(r"<[^>]+>", " ", content)
    files = [a.get_filename() for a in msg.iter_attachments()]
    att = f"<p><b>{h(t('pv_attachments'))}:</b> {h(', '.join(x for x in files if x))}</p>" if files else ""
    return f"<table>{head}</table>{att}<pre>{h(content)}</pre>"

def zip_html(p):
    with zipfile.ZipFile(p) as z:
        items = z.infolist()
        rows = "".join(f"<tr><td>{h(i.filename)}</td><td>{i.file_size:,}</td><td>{datetime(*i.date_time):%Y-%m-%d %H:%M}</td></tr>" for i in items[:3000])
    return f"<p>{h(t('pv_zip_count', n=len(items)))}</p><table><tr><th>{h(t('col_name'))}</th><th>{h(t('col_size_bytes'))}</th><th>{h(t('col_date'))}</th></tr>{rows}</table>"

def looks_like_text(p):
    with p.open("rb") as f: chunk = f.read(4096)
    if b"\x00" in chunk: return False
    try:
        chunk.decode("utf-8"); return True
    except UnicodeDecodeError:
        return sum(32 <= b < 127 or b in (9, 10, 13) or b >= 192 for b in chunk) / max(len(chunk), 1) > 0.9

def unsupported(p, message):
    return {"kind": "unsupported", "name": p.name, "message": message}

def preview_payload(p: Path, raw_url: str):
    """Описание предпросмотра для файла p; raw_url отдаёт сам файл (каталог или карантин заявок)."""
    ext = p.suffix.lower()
    try:
        if ext == ".pdf": return {"kind": "pdf", "url": raw_url, "name": p.name}
        if ext in IMAGE_EXTS: return {"kind": "image", "url": raw_url, "name": p.name}
        if ext in IMAGE_CONVERT:
            if Image is None: return unsupported(p, t("pv_need_pillow"))
            return converted(p, "png", "image", convert_image, "pv_preparing_image")
        if ext in VIDEO_EXTS: return {"kind": "video", "url": raw_url, "name": p.name}
        if ext in AUDIO_EXTS: return {"kind": "audio", "url": raw_url, "name": p.name}
        if ext in VIDEO_CONVERT | AUDIO_CONVERT:
            if not FFMPEG: return unsupported(p, t("pv_need_ffmpeg"))
            video = ext in VIDEO_CONVERT
            return converted(p, "mp4" if video else "mp3", "video" if video else "audio", convert_media, "pv_converting_video" if video else "pv_converting_audio")
        if ext in {".xlsx", ".xlsm"}: return {"kind": "html", "html": xlsx_html(p), "name": p.name}
        if ext == ".csv": return {"kind": "html", "html": csv_html(p), "name": p.name}
        if ext in OFFICE_EXTS:
            if SOFFICE: return converted(p, "pdf", "pdf", convert_office, "pv_converting_doc")
            if ext in {".docx", ".docm"}: return {"kind": "html", "html": docx_html(p), "name": p.name, "note": t("pv_simple_docx")}
            if ext in {".pptx", ".ppsx", ".pptm"}: return {"kind": "html", "html": pptx_html(p), "name": p.name, "note": t("pv_simple_pptx")}
            return unsupported(p, t("pv_need_lo"))
        if ext == ".eml": return {"kind": "html", "html": eml_html(p), "name": p.name}
        if ext == ".zip": return {"kind": "html", "html": zip_html(p), "name": p.name}
        if ext in HTML_EXTS: return {"kind": "frame", "url": raw_url, "name": p.name}
        if ext in TEXT_EXTS or (ext not in {".dwg", ".msg", ".exe", ".dll", ".msi"} and looks_like_text(p)):
            with p.open("rb") as f: raw = f.read(1_000_001)
            return {"kind": "text", "name": p.name, "content": raw[:1_000_000].decode("utf-8-sig", errors="replace"), "truncated": len(raw) > 1_000_000}
    except Exception as e:
        return {"kind": "error", "name": p.name, "message": t("pv_read_error", error=e)}
    hint = {".dwg": "pv_hint_dwg", ".msg": "pv_hint_msg"}.get(ext, "pv_hint_other")
    if ext in EXEC_EXTS: hint = "pv_hint_exec"
    return unsupported(p, t(hint))
