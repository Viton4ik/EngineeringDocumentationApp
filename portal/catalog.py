"""Сканирование каталога, сопоставление Siemens и распознавание названий по первой странице."""
from __future__ import annotations
import os, re, logging, threading
from pathlib import Path
from datetime import datetime
from .config import DEFAULT_ROOT, NAME_FILE, NUMBER_FILE, DOC_SHARE, SHARE_NAME, HOST, EXCLUDED, TITLE_SCAN, NAME_FROM_FILENAME, file_kind
from .db import execute, query

logging.getLogger("pypdf").setLevel(logging.ERROR)  # битые PDF не должны засорять консоль сервера

CACHE = {"files": [], "folders": [], "error": None, "scanned_at": None, "siemens_count": 0}
SKIP_FILES = {"thumbs.db", "desktop.ini", ".ds_store"}
TITLES = {"running": False, "done": 0, "total": 0, "found": 0, "log": []}
LOCK = threading.RLock()

def read_lines(path: Path):
    try:
        return [line.strip() for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines()]
    except OSError:
        return []

def siemens_map():
    names, numbers = read_lines(NAME_FILE), read_lines(NUMBER_FILE)
    # Same line-to-line mapping as directory.py; first matching number wins.
    result = {}
    for number, name in zip(numbers, names):
        number, name = number.strip(), name.strip()
        if number and name and number not in result:
            result[number] = name
    return result

def root():
    return DEFAULT_ROOT.resolve()

def make_unc(path: Path, base: Path = None):
    base = base or root()
    if DOC_SHARE:  # путь для сервера и выгрузок; в браузере имя сервера подставляется из адресной строки
        prefix = DOC_SHARE if DOC_SHARE.startswith("\\\\") else "\\\\" + HOST + "\\" + SHARE_NAME
        return prefix + "\\" + str(path.relative_to(base)).replace("/", "\\")
    absolute = str(path)
    if len(absolute) > 2 and absolute[1] == ":":
        return "\\\\" + HOST + "\\" + absolute[0] + "$" + absolute[2:].replace("/", "\\")
    return "\\\\" + HOST + "\\" + absolute.replace("/", "\\").lstrip("\\")

def load_titles():
    try: return {r["path"]: r for r in query("SELECT * FROM titles")}
    except Exception: return {}

def load_meta():
    try: return {r["path"]: r for r in query("SELECT * FROM doc_meta")}
    except Exception: return {}

def make_record(path: Path, base: Path, mapping, stored, meta):
    st = path.stat()
    parts = path.parent.relative_to(base).parts
    absolute = str(path.resolve())
    filename = path.name
    doc_name, source, description = filename, "file", ""
    m = meta.get(absolute)
    by_filename = bool(parts) and parts[0].casefold() in NAME_FROM_FILENAME
    if m:  # одобренные администратором документы: название и описание из заявки
        doc_name, source, description = m["title"] or filename, "approved", m["description"] or ""
    else:
        for number, name in mapping.items():
            if number in filename[:26]:
                doc_name, source = name, "siemens"
                break
        if source == "file" and by_filename:  # например Check lists: название = имя файла, дерево = папки
            doc_name, source = re.sub(r"[_\s]+", " ", path.stem).strip() or filename, "filename"
        if source == "file":
            t = stored.get(absolute)
            if t and t["mtime"] == st.st_mtime and t["size"] == st.st_size and t["title"]:
                doc_name, source = t["title"], "content"
    ext = path.suffix.lower()
    return {"id": absolute, "name": filename, "doc_name": doc_name, "title_source": source, "description": description,
            "documentation": parts[0] if len(parts) >= 1 else base.name, "rfq": parts[1] if len(parts) >= 2 else "",
            "system": parts[2] if len(parts) >= 3 else "", "unit": parts[3] if len(parts) >= 4 else (parts[-1] if parts else base.name),
            "folder": str(path.parent), "relative": str(path.relative_to(base)), "path": absolute, "unc": make_unc(path, base),
            "extension": ext, "kind": file_kind(ext), "size": st.st_size, "mtime": st.st_mtime,
            "modified": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="minutes")}

def scan_documents():
    base = root()
    mapping, stored, meta = siemens_map(), load_titles(), load_meta()
    records, folders, error = [], [], None
    if not base.exists() or not base.is_dir():
        error = f"Каталог не найден или недоступен: {base}"
    else:
        for current, dirs, files in os.walk(base, onerror=lambda e: None):
            dirs[:] = sorted((d for d in dirs if d not in EXCLUDED and not d.startswith(".")), key=str.casefold)
            rel = Path(current).relative_to(base)
            if rel.parts: folders.append("/".join(rel.parts))  # и пустые папки: они нужны для выбора раздела
            for filename in sorted(files, key=str.casefold):
                if filename.startswith("~$") or filename.casefold() in SKIP_FILES: continue  # временные и служебные файлы
                try: records.append(make_record(Path(current) / filename, base, mapping, stored, meta))
                except OSError: continue
    with LOCK:
        CACHE.update(files=records, folders=folders, error=error, scanned_at=datetime.now().isoformat(timespec="seconds"), siemens_count=len(mapping))
    if TITLE_SCAN: start_title_scan()
    return records, error

def add_files(paths):
    """Добавить в кэш только что одобренные файлы без полного пересканирования."""
    base = root()
    mapping, stored, meta = siemens_map(), load_titles(), load_meta()
    with LOCK:
        for p in paths:
            rec = make_record(Path(p), base, mapping, stored, meta)
            add_folder(str(Path(p).parent.relative_to(base)))
            CACHE["files"] = [r for r in CACHE["files"] if r["id"] != rec["id"]] + [rec]
            TITLES["log"].append([rec["id"], rec["doc_name"], rec["title_source"], rec])

def add_folder(rel):
    """Новая папка (из одобренной заявки или переноса) сразу появляется в дереве выбора."""
    parts = [p for p in rel.replace("\\", "/").split("/") if p]
    with LOCK:
        for i in range(1, len(parts) + 1):
            f = "/".join(parts[:i])
            if f not in CACHE["folders"]: CACHE["folders"].append(f)

def move_record(old_path, new_path):
    """Обновить кэш после переноса: клиенты получат запись 'moved' и перенесут избранное."""
    base = root()
    mapping, stored, meta = siemens_map(), load_titles(), load_meta()
    rec = make_record(Path(new_path), base, mapping, stored, meta)
    add_folder(str(Path(new_path).parent.relative_to(base)))
    with LOCK:
        CACHE["files"] = [r for r in CACHE["files"] if r["id"] not in (old_path, rec["id"])] + [rec]
        TITLES["log"].append([old_path, rec["doc_name"], "moved", rec])
    return rec

def find_record(path):
    target = str(Path(path))
    with LOCK: return next((r for r in CACHE["files"] if r["path"] == target), None)

# ---------- Распознавание названий ----------
BAD_LINE = re.compile(r"^(page|стр\.?|страница|лист|sheet|strana|rev\.?|revision|ред\.?|date|дата|datum|copyright|©|www\.|http|tel|fax|e-?mail|confidential|"
                      r"конфиденциально|microsoft (word|excel|powerpoint)|untitled|без названия|документ\d*|document\d*|slide|презентация\d*)\b", re.I)

def good_title(s, filename=""):
    s = re.sub(r"\s+", " ", str(s or "")).strip(" .:;,-_|")
    if not (6 <= len(s) <= 180) or BAD_LINE.match(s): return ""
    letters = sum(ch.isalpha() for ch in s)
    if letters < 4 or letters / len(s) < 0.45: return ""
    if Path(filename).stem.casefold() == s.casefold() or s.lower().endswith((".doc", ".docx", ".pdf", ".xls", ".xlsx", ".dwg")): return ""
    return s[:150]

def first_good(lines, filename):
    for line in lines:
        t = good_title(line, filename)
        if t: return t
    return ""

def title_from_pdf(path: Path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try: reader.decrypt("")
        except Exception: return ""
    meta = reader.metadata
    t = good_title(meta.title if meta else "", path.name)
    if t: return t
    if not reader.pages: return ""
    return first_good((reader.pages[0].extract_text() or "").splitlines(), path.name)

def title_from_docx(path: Path):
    import docx
    d = docx.Document(str(path))
    t = good_title(d.core_properties.title, path.name)
    if t: return t
    paras = [p for p in d.paragraphs[:80] if p.text.strip()]
    styled = [p.text for p in paras if (p.style is not None and p.style.name or "").lower().startswith(("title", "heading", "заголовок", "название"))]
    cells = [c.text for tbl in d.tables[:2] for row in tbl.rows[:15] for c in row.cells]
    return first_good(styled, path.name) or first_good([p.text for p in paras], path.name) or first_good(cells, path.name)

def title_from_pptx(path: Path):
    from pptx import Presentation
    prs = Presentation(str(path))
    t = good_title(prs.core_properties.title, path.name)
    if t or not prs.slides: return t
    slide = prs.slides[0]
    if slide.shapes.title is not None:
        t = good_title(slide.shapes.title.text, path.name)
        if t: return t
    return first_good([sh.text_frame.text for sh in slide.shapes if sh.has_text_frame], path.name)

def title_from_xlsx(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        t = good_title(wb.properties.title, path.name)
        if t: return t
        ws = wb.worksheets[0]
        return first_good([v for row in ws.iter_rows(max_row=30, max_col=15, values_only=True) for v in row if isinstance(v, str)], path.name)
    finally:
        wb.close()

def title_from_text(path: Path):
    with path.open("rb") as f: raw = f.read(8192)
    return first_good(raw.decode("utf-8", errors="replace").splitlines(), path.name)

TITLE_READERS = {".pdf": title_from_pdf, ".docx": title_from_docx, ".docm": title_from_docx, ".pptx": title_from_pptx,
                 ".ppsx": title_from_pptx, ".xlsx": title_from_xlsx, ".xlsm": title_from_xlsx, ".txt": title_from_text, ".md": title_from_text}

def extract_title(path: Path):
    reader = TITLE_READERS.get(path.suffix.lower())
    if not reader: return ""
    try: return reader(path)
    except Exception: return ""

def save_title(rec, title, source="content"):
    execute("INSERT OR REPLACE INTO titles(path,mtime,size,title,source) VALUES(?,?,?,?,?)", (rec["path"], rec["mtime"], rec["size"], title, source))
    if title:
        with LOCK:
            rec["doc_name"], rec["title_source"] = title, "content"
            TITLES["found"] += 1
            TITLES["log"].append([rec["id"], title, "content"])

def start_title_scan():
    with LOCK:
        if TITLES["running"]: return
        stored = load_titles()
        todo = [r for r in CACHE["files"] if r["title_source"] == "file" and r["extension"] in TITLE_READERS
                and not (r["path"] in stored and stored[r["path"]]["mtime"] == r["mtime"] and stored[r["path"]]["size"] == r["size"])]
        TITLES.update(running=bool(todo), done=0, total=len(todo), found=0)
    if not todo: return
    def work():
        print(f"Распознавание названий: {len(todo)} файлов в фоне…")
        for rec in todo:
            save_title(rec, extract_title(Path(rec["path"])))
            with LOCK: TITLES["done"] += 1
        with LOCK: TITLES["running"] = False
        print(f"Распознавание названий завершено: найдено {TITLES['found']} из {len(todo)}")
    threading.Thread(target=work, daemon=True).start()
